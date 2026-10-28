"""Read local exported attachments. Never opens WXS, executes macros, or follows document links."""
import argparse
import csv
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile

HERE = Path(__file__).resolve().parent
NARAE = HERE / 'readers/narae_hangul_office/narae.py'


def json_file(path, value):
    with path.open('x', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, default=str)


def render_pdf(source, directory):
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument(str(source))
    try:
        paths = []
        for index in range(len(document)):
            page = document[index]
            try:
                bitmap = page.render(scale=1.5)
                path = directory / f'page-{index + 1:04d}.png'
                bitmap.to_pil().save(path)
                bitmap.close()
                paths.append(str(path))
            finally:
                page.close()
        return paths
    finally:
        document.close()


def ocr_images(paths, directory):
    results, warnings = [], []
    if os.name != 'nt':
        return results, ['Windows Korean OCR unavailable on this platform; visual review required']
    from PIL import Image
    for i, path in enumerate(paths):
        preview = directory / f'ocr-input-{i + 1:04d}.png'
        with Image.open(path) as image:
            image = image.convert('RGB')
            image.thumbnail((2000, 2000))
            image.save(preview)
        output = directory / f'ocr-{i + 1:04d}.json'
        process = subprocess.run(['powershell.exe', '-NoProfile', '-NonInteractive', '-File', str(HERE / 'readers/image-ocr.ps1'),
                                  '-Path', str(preview), '-Out', str(output)], capture_output=True, text=True, encoding='utf-8', errors='replace',
                                 timeout=90, creationflags=subprocess.CREATE_NO_WINDOW)
        if process.returncode == 0 and output.is_file():
            results.append(json.loads(output.read_text(encoding='utf-8-sig')))
        else:
            warnings.append(f'OCR failed for image {i + 1}; visual review required')
    return results, warnings


def native_hangul(source, directory, render):
    job = directory / 'narae-job.json'
    raw = directory / 'narae-result.json'
    json_file(job, {'schema': 1, 'action': 'inspect', 'input': str(source)})
    options = {'capture_output': True, 'text': True, 'encoding': 'utf-8', 'errors': 'replace', 'timeout': 300}
    if os.name == 'nt':
        options['creationflags'] = subprocess.CREATE_NO_WINDOW
    process = subprocess.run([sys.executable, '-X', 'utf8', str(NARAE), 'run', str(job), '--result', str(raw)], **options)
    if not raw.is_file():
        raise ValueError('Narae Hangul returned no result: ' + process.stderr[-1000:])
    result = json.loads(raw.read_text(encoding='utf-8-sig'))
    if process.returncode or not result.get('success') or not result.get('input_unchanged'):
        raise ValueError('Narae Hangul refused/failed: ' + str(result.get('error', process.stderr[-1000:])))
    if result.get('textDetailsLoaded') is False:
        raise ValueError('변경 추적 문서는 나래 한글 오피스 live-read로 현재 본문을 확인하세요. 빈 문서로 처리하지 않습니다.')
    paragraphs = [{'id': p['id'], 'text': p.get('text', '')} for p in result.get('paragraphs', [])]
    text = '\n'.join(p['text'] for p in paragraphs)
    images = result.get('images', [])
    view = []
    if render or (not text.strip() and images):
        pdf = directory / 'narae-view.pdf'
        export_result = directory / 'narae-export.json'
        process = subprocess.run([sys.executable, '-X', 'utf8', str(NARAE), 'export', str(source), '--pdf', str(pdf), '--result', str(export_result)], **options)
        if process.returncode or not pdf.is_file():
            raise ValueError('Narae Hangul preview failed: ' + process.stderr[-1000:])
        view = render_pdf(pdf, directory)
    return {'reader': 'Narae Hangul Office CLI', 'text': text, 'paragraphs': paragraphs,
            'tables': result.get('tables', []), 'embedded_images': images, 'visual_pages': view,
            'needs_visual_review': bool(images), 'warnings': result.get('warnings', []),
            'read_restrictions': result.get('read_only'), 'source_unchanged': True}


def embedded_images(source, directory, prefix):
    paths = []
    with zipfile.ZipFile(source) as archive:
        for info in archive.infolist():
            if info.filename.startswith(prefix) and not info.is_dir():
                if info.file_size > 100_000_000:
                    raise ValueError('Embedded image exceeds 100 MB limit')
                path = directory / f'image-{len(paths) + 1:04d}{Path(info.filename).suffix}'
                with path.open('xb') as stream:
                    stream.write(archive.read(info))
                paths.append(str(path))
    return paths


def decode_text(data):
    if data.startswith((b'\xff\xfe', b'\xfe\xff')):
        return data.decode('utf-16'), 'utf-16'
    for encoding in ('utf-8-sig', 'cp949'):
        try:
            return data.decode(encoding), encoding
        except UnicodeError:
            pass
    raise ValueError('Unknown text encoding; no lossy decoding used')


def word_container_text(container):
    lines = []
    for block in container.iter_inner_content():
        if hasattr(block, 'rows'):
            for row in block.rows:
                lines.extend(word_container_text(cell) for cell in row.cells)
        else:
            lines.append(block.text)
    return '\n'.join(lines)


def slide_shapes(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, 'shapes'):
            yield from slide_shapes(shape.shapes)


def extract(source, directory, render=False):
    ext = source.suffix.lower()
    if ext in ('.hwp', '.hwpx'):
        return native_hangul(source, directory, render)
    if ext == '.pdf':
        from pypdf import PdfReader
        pdf = PdfReader(source)
        from pypdf.errors import FileNotDecryptedError
        try:
            # PdfReader already opens PDFs whose user password is empty. Do not
            # reject readable owner-protected files or try any supplied password.
            pages = [{'page': i + 1, 'text': p.extract_text() or ''} for i, p in enumerate(pdf.pages)]
        except FileNotDecryptedError as error:
            raise ValueError('PDF requires a password; no password guessing') from error
        scanned = any(len(p['text'].strip()) < 200 for p in pages)
        images = render_pdf(source, directory) if render or scanned else []
        ocr, warnings = ocr_images(images, directory) if scanned else ([], [])
        return {'reader': 'pypdf', 'pages': pages, 'text': '\n\n'.join(p['text'] for p in pages),
                'ocr_text': '\n'.join(p['text'] for p in ocr), 'ocr': ocr, 'warnings': warnings,
                'visual_pages': images, 'needs_visual_review': scanned, 'page_count': len(pages),
                'encrypted': pdf.is_encrypted}
    if ext in ('.xlsx', '.xlsm', '.xltx', '.xltm'):
        import openpyxl
        book = openpyxl.load_workbook(source, data_only=False, read_only=True, keep_links=False)
        cached = openpyxl.load_workbook(source, data_only=True, read_only=True, keep_links=False)
        try:
            sheets, cells = [], 0
            for sheet in book:
                rows = []
                for row in sheet.iter_rows():
                    values = []
                    for cell in row:
                        if cell.value is not None:
                            cells += 1
                            if cells > 1_000_000:
                                raise ValueError('Workbook exceeds one million nonempty cells; not silently truncated')
                            item = {'cell': cell.coordinate, 'value': cell.value}
                            if cell.data_type == 'f':
                                item['cached_value'] = cached[sheet.title][cell.coordinate].value
                            values.append(item)
                    if values:
                        rows.append(values)
                sheets.append({'name': sheet.title, 'state': sheet.sheet_state, 'rows': rows})
            images = embedded_images(source, directory, 'xl/media/')
            return {'reader': 'openpyxl', 'sheets': sheets, 'text': '\n'.join(s['name'] + '\n' + '\n'.join(' | '.join(str(c['value']) for c in r) for r in s['rows']) for s in sheets),
                    'visual_pages': images, 'needs_visual_review': bool(images), 'formulas_executed': False}
        finally:
            book.close(); cached.close()
    if ext == '.xls':
        import xlrd
        book = xlrd.open_workbook(source)
        sheets = [{'name': s.name, 'rows': [s.row_values(i) for i in range(s.nrows)]} for s in book.sheets()]
        return {'reader': 'xlrd', 'sheets': sheets, 'text': '\n'.join(s['name'] + '\n' + '\n'.join(' | '.join(map(str, r)) for r in s['rows']) for s in sheets), 'formulas_executed': False}
    if ext == '.docx':
        from docx import Document
        document = Document(source)
        paragraphs = [p.text for p in document.paragraphs]
        tables = [[[' | '.join(c.text for c in r.cells)] for r in t.rows] for t in document.tables]
        sections = [{name: word_container_text(getattr(section, name))
                     for name in ('header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer')}
                    for section in document.sections]
        images = embedded_images(source, directory, 'word/media/')
        return {'reader': 'python-docx', 'paragraphs': paragraphs, 'tables': tables, 'sections': sections,
                'text': word_container_text(document) + '\n' + '\n'.join('\n'.join(section.values()) for section in sections),
                'visual_pages': images, 'needs_visual_review': bool(images)}
    if ext == '.pptx':
        from pptx import Presentation
        deck = Presentation(source)
        slides = []
        for i, slide in enumerate(deck.slides):
            texts = []
            for shape in slide_shapes(slide.shapes):
                if shape.has_text_frame:
                    texts.append(shape.text)
                if shape.has_table:
                    texts.extend(' | '.join(c.text for c in r.cells) for r in shape.table.rows)
            if slide.has_notes_slide:
                texts.append(slide.notes_slide.notes_text_frame.text)
            slides.append({'slide': i + 1, 'text': '\n'.join(texts)})
        images = embedded_images(source, directory, 'ppt/media/')
        return {'reader': 'python-pptx', 'slides': slides, 'text': '\n\n'.join(s['text'] for s in slides),
                'visual_pages': images, 'needs_visual_review': bool(images)}
    if ext in ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tif', '.tiff', '.webp'):
        from PIL import Image, ImageOps, ImageSequence
        paths = []
        with Image.open(source) as image:
            original_size = list(image.size)
            if image.format == 'JPEG' and image.width * image.height > 100_000_000:
                image.draft('RGB', (4000, 4000))
            for i, frame in enumerate(ImageSequence.Iterator(image)):
                if frame.width * frame.height > 100_000_000:
                    raise ValueError('Image exceeds 100 million pixels')
                path = directory / f'frame-{i + 1:04d}.png'
                ImageOps.exif_transpose(frame).convert('RGB').save(path)
                paths.append(str(path))
        ocr, warnings = ocr_images(paths, directory)
        return {'reader': 'Pillow / Windows.Media.Ocr', 'text': '\n'.join(p['text'] for p in ocr), 'ocr': ocr, 'warnings': warnings,
                'original_size': original_size,
                'visual_pages': paths, 'needs_visual_review': True,
                'coverage': 'Images prepared for visual reading; filenames/dimensions are not content review'}
    if ext in ('.txt', '.csv', '.tsv', '.md', '.log', '.json', '.xml', '.html', '.htm'):
        text, encoding = decode_text(source.read_bytes())
        result = {'reader': 'text', 'text': text, 'encoding': encoding}
        if ext in ('.csv', '.tsv'):
            result['rows'] = list(csv.reader(io.StringIO(text), delimiter='\t' if ext == '.tsv' else ','))
        return result
    raise ValueError('Unsupported attachment format: ' + ext + '; do not mark as read')


def read_file(source, directory, render=False):
    source = source.resolve(strict=True)
    if not source.is_file():
        raise ValueError('Input must be a file')
    directory.mkdir(parents=True, exist_ok=False)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    result = {'source': str(source), 'sha256': digest, 'status': 'failed'}
    try:
        result.update(extract(source, directory, render))
        if hashlib.sha256(source.read_bytes()).hexdigest() != digest:
            raise ValueError('Source changed while reading')
        result.update(status='extracted', source_unchanged=True, content_reviewed=False)
    except Exception as error:
        result['error'] = str(error)
    json_file(directory / 'result.json', result)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path, nargs='?')
    parser.add_argument('--out-dir', type=Path)
    parser.add_argument('--render', action='store_true', help='Render all PDF/Hangul pages for visual review')
    parser.add_argument('--doctor', action='store_true')
    args = parser.parse_args()
    if args.doctor:
        probe = subprocess.run([sys.executable, '-X', 'utf8', str(NARAE), 'doctor'],
                               capture_output=True, encoding='utf-8', errors='replace', timeout=30,
                               creationflags=0x08000000 if os.name == 'nt' else 0)
        try:
            office = json.loads(probe.stdout)
        except ValueError:
            office = {'success': False, 'error': probe.stderr[-1000:]}
        result = {'narae_cli': NARAE.is_file(), 'dependencies': {m: importlib.util.find_spec(m) is not None for m in ('pypdf', 'pypdfium2', 'openpyxl', 'xlrd', 'docx', 'pptx', 'PIL')},
                  'narae_office': office,
                  'korean_ocr': 'Windows language pack checked when reading an image' if os.name == 'nt' else 'unavailable; visual review required'}
    else:
        if not args.input or not args.out_dir:
            parser.error('input and --out-dir are required')
        result = read_file(args.input, args.out_dir, args.render)
        result = {k: result[k] for k in ('source', 'status', 'error', 'needs_visual_review') if k in result}
        result['result'] = str((args.out_dir / 'result.json').resolve())
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return int(result.get('status') == 'failed' or (args.doctor and (not office.get('success') or not all(result['dependencies'].values()))))


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    raise SystemExit(main())
