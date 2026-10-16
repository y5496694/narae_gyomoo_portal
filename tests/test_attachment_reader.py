"""Actual file round trips for the bundled readers; no application login or user files."""
from pathlib import Path
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

SCRIPTS = Path(__file__).resolve().parents[1] / 'skills/wxs-cli/scripts'
sys.path.insert(0, str(SCRIPTS))
from attachment_reader import NARAE, decode_text, read_file
from wxs_read import safe_attachment_names


class AttachmentReaderTests(unittest.TestCase):
    def test_korean_cp949_is_not_misread_as_bomless_utf16(self):
        text = '한글 신청 안내'
        self.assertEqual(decode_text(text.encode('cp949')), (text, 'cp949'))
        self.assertEqual(decode_text(text.encode('utf-16')), (text, 'utf-16'))

    def test_native_hwp_and_hwpx_use_bundled_narae_and_preserve_source(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for ext in ('hwp', 'hwpx'):
                with self.subTest(format=ext):
                    source = root/f'example.{ext}'
                    job = root/f'create-{ext}.json'
                    job.write_text(json.dumps({'schema':1, 'action':'create', 'format':ext, 'output':str(source),
                                               'blocks':[{'type':'paragraph','text':'예시 문서의 제출 기한은 10월 21일입니다.'}]}, ensure_ascii=False), encoding='utf-8')
                    process = subprocess.run([sys.executable, '-X', 'utf8', str(NARAE), 'run', str(job)], capture_output=True, text=True, encoding='utf-8', timeout=120)
                    self.assertEqual(process.returncode, 0, process.stderr)
                    before = hashlib.sha256(source.read_bytes()).hexdigest()
                    result = read_file(source, root/f'read-{ext}')
                    self.assertEqual(result['status'], 'extracted', result.get('error'))
                    self.assertEqual(result['reader'], 'Narae Hangul (bundled CLI)')
                    self.assertIn('제출 기한은 10월 21일', result['text'])
                    self.assertEqual(before, hashlib.sha256(source.read_bytes()).hexdigest())

    def test_pdf_text_and_rendered_page_both_available(self):
        from pypdf import PdfWriter
        from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'example.pdf';writer=PdfWriter()
            page=writer.add_blank_page(width=300,height=100)
            font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
            page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
            stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 10 40 Td (Application deadline 2026-10-21 18:00) Tj ET')
            page[NameObject('/Contents')]=writer._add_object(stream)
            writer.write(source)
            result=read_file(source,root/'out',render=True)
            self.assertEqual(result['status'],'extracted',result.get('error'))
            self.assertIn('2026-10-21 18:00',result['text'])
            self.assertEqual(result['page_count'],1);self.assertTrue(Path(result['visual_pages'][0]).is_file())

    def test_path_traversal_reserved_names_and_case_collisions_rejected(self):
        for names in [['../other.hwp'], ['C:\\other.hwp'], ['CON.txt'], ['a.xlsx', 'A.XLSX']]:
            with self.subTest(names=names), self.assertRaises(ValueError):
                safe_attachment_names([{'name': n} for n in names])

    def test_pdf_open_password_is_not_guessed_and_empty_user_password_is_readable(self):
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for password in ('', 'private-example'):
                writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
                writer.encrypt(user_password=password, owner_password='owner-example')
                source = root / ('open.pdf' if not password else 'locked.pdf'); writer.write(source)
                with patch('attachment_reader.render_pdf', return_value=[]):
                    result = read_file(source, root / source.stem)
                self.assertEqual(result['status'], 'failed' if password else 'extracted')
                if password:
                    self.assertIn('requires a password', result['error'])

    def test_hidden_sheet_formulas_and_all_rows_are_preserved_without_execution(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); source = root/'test.xlsx'
            book = openpyxl.Workbook(); book.active['A1'] = '신청 안내'; book.active['B2'] = '=1+2'
            hidden = book.create_sheet('숨김 서식'); hidden.sheet_state = 'hidden'; hidden['C5'] = '제출 기한'
            book.save(source); book.close()
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            result = read_file(source, root/'out')
            self.assertEqual(result['status'], 'extracted')
            self.assertEqual([s['name'] for s in result['sheets']], ['Sheet', '숨김 서식'])
            self.assertEqual(result['sheets'][1]['state'], 'hidden')
            self.assertIn('제출 기한', result['text'])
            self.assertIn('=1+2', result['text'])
            self.assertFalse(result['formulas_executed'])
            self.assertEqual(before, hashlib.sha256(source.read_bytes()).hexdigest())
            self.assertFalse(result['content_reviewed'])

    def test_word_paragraph_table_and_footer(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp); source=root/'test.docx'
            doc=Document();doc.add_paragraph('본문 안내');doc.add_table(rows=1,cols=1).cell(0,0).text='첨부 서식'
            doc.sections[0].footer.paragraphs[0].text='하단 주의사항'
            doc.sections[0].first_page_header.paragraphs[0].text='첫 쪽 머리말'
            doc.tables[0].cell(0,0).add_table(rows=1,cols=1).cell(0,0).text='중첩 표 내용'
            doc.save(source)
            result=read_file(source,root/'out')
            self.assertEqual(result['status'],'extracted')
            self.assertIn('본문 안내',result['text']);self.assertIn('첨부 서식',result['text'])
            self.assertIn('하단 주의사항',result['sections'][0]['footer'])
            self.assertIn('첫 쪽 머리말',result['text']);self.assertIn('중첩 표 내용',result['text'])

    def test_presentation_all_slides_and_notes(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'test.pptx';deck=Presentation()
            for label in ['첫째 내용','둘째 내용']:
                slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text=label
                slide.notes_slide.notes_text_frame.text='강사 노트'
                group=slide.shapes.add_group_shape(); group.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(1)).text='그룹 안 내용'
            deck.save(source);result=read_file(source,root/'out')
            self.assertEqual(result['status'],'extracted');self.assertEqual(len(result['slides']),2)
            self.assertIn('둘째 내용',result['text']);self.assertIn('강사 노트',result['text'])
            self.assertIn('그룹 안 내용',result['text'])

    def test_image_frames_require_review_and_unsupported_input_never_succeeds(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'test.gif'
            Image.new('RGB',(10,10),'white').save(source,save_all=True,append_images=[Image.new('RGB',(10,10),'black')])
            with patch('attachment_reader.ocr_images', return_value=([],['test without OCR'])):
                result=read_file(source,root/'image')
            self.assertEqual(len(result['visual_pages']),2);self.assertTrue(result['needs_visual_review'])
            unknown=root/'test.exe';unknown.write_bytes(b'not a document')
            result=read_file(unknown,root/'unknown')
            self.assertEqual(result['status'],'failed');self.assertNotIn('content_reviewed',result)


if __name__ == '__main__': unittest.main()
